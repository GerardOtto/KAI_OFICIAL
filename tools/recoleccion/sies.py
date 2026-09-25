# -*- coding: utf-8 -*-
"""T2.1 — SIES / mifuturo.cl: personal académico, matrícula y titulados.

Es la tarea más importante del plan: de aquí salen los denominadores reales de
media docena de métricas de THE y QS —estudiantes, académicos JCE, académicos con
doctorado, extranjeros y titulados por nivel— para **todas** las universidades
chilenas y en serie anual, no solo para las treinta que aparecen en los rankings.

Las tres bases históricas del SIES, localizadas en mifuturo.cl:

    Personal académico 2008-2025   PAC_web_2008_2025_SIES_EE.xlsx      2 MB
    Titulados 2007-2025            TITULADO_2007-2025_web_…xlsx       41 MB
    Matrícula 2007-2026            Matricula_2007_2026_WEB_…zip       11 MB

Cómo vienen: la de personal académico trae **una fila por institución y año**, con
la cabecera repartida en tres filas —grupo, subgrupo y columna— y dos hojas
paralelas, una con el número de personas y otra con las jornadas completas
equivalentes. Las de titulados y matrícula son microdatos por programa y hay que
agregarlos.

Cada variable lleva en `definicion` cómo se calculó, porque la cifra depende de esa
decisión: «académicos JCE» es la suma de jornadas equivalentes y no un recuento de
personas, y esa diferencia cambia el resultado de THE Student-staff ratio.

Uso:
    python tools/recoleccion/sies.py                 # descarga y procesa
    python tools/recoleccion/sies.py --solo-procesar
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import navegador as nav  # noqa: E402

FUENTE = "sies"

BASES = {
    "personal_academico": ("https://mifuturo.cl/wp-content/uploads/2025/12/PAC_web_2008_2025_SIES_EE.xlsx",
                           "personal_academico_2008_2025.xlsx"),
    "titulados": ("https://www.mifuturo.cl/wp-content/uploads/2026/05/TITULADO_2007-2025_web_27_05_2026_E-1.xlsx",
                  "titulados_2007_2025.xlsx"),
    "matricula": ("https://www.mifuturo.cl/wp-content/uploads/2026/07/Matricula_2007_2026_WEB_10_07_2026.zip",
                  "matricula_2007_2026.zip"),
}

DESDE = 2015  # el plan pide desde 2015: antes la serie pesa más y aporta menos


def plano(texto) -> str:
    """Texto sin tildes, en minúsculas y con los espacios colapsados."""
    sin_tildes = "".join(c for c in unicodedata.normalize("NFD", str(texto or ""))
                         if unicodedata.category(c) != "Mn")
    return " ".join(sin_tildes.lower().split())


def numero(valor) -> float | None:
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto or texto in {"-", "s/i", "sin informacion"}:
        return None
    # Los CSV del SIES usan coma decimal; los XLSX ya traen números.
    texto = texto.replace(".", "").replace(",", ".") if texto.count(",") == 1 else texto
    try:
        return float(texto)
    except ValueError:
        return None


def descargar() -> None:
    base = nav.carpeta(FUENTE) / "raw"
    for nombre, (url, archivo) in BASES.items():
        ruta = nav.descargar(url, base / archivo, FUENTE, notas=nombre)
        print(f"  {nombre:20} {f'{ruta.stat().st_size / 1e6:.1f} MB' if ruta else 'sin datos'}")


# --- Personal académico -------------------------------------------------------

def _columnas_compuestas(ws) -> list[str]:
    """Nombre completo de cada columna, uniendo las tres filas de cabecera.

    La primera fila trae el grupo («N° de JCE por nivel de formación») en una celda
    combinada, así que su valor se arrastra hacia la derecha hasta que aparece otro.
    """
    f1, f2, f3 = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
    nombres, grupo, subgrupo = [], "", ""
    for a, b, c in zip(f1, f2, f3):
        if a:
            grupo, subgrupo = plano(a), ""
        if b:
            subgrupo = plano(b)
        nombres.append(" | ".join(x for x in (grupo, subgrupo, plano(c)) if x))
    return nombres


def _indice(columnas: list[str], grupo: str, hoja: str) -> int | None:
    for i, nombre in enumerate(columnas):
        if grupo in nombre and nombre.endswith(hoja):
            return i
    return None


def personal_academico(ruta: Path, resolver) -> tuple[list[dict], set[str]]:
    """Académicos por institución y año: total, JCE, doctorado y extranjeros."""
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    # Las dos hojas tienen la misma forma; cambia la unidad de lo que cuentan.
    HOJAS = {
        "BD_Acádemicos_Número": ("personas", {
            "academicos_total": ("por insitucion", "total general",
                                 "Personas con función académica declaradas por la institución"),
            "academicos_doctorado": ("por nivel de formacion", "doctor",
                                     "Personas cuyo grado académico máximo es doctor"),
            "academicos_extranjeros": ("por nacionalidad", "extranjero",
                                       "Personas de nacionalidad distinta de la chilena"),
        }),
        "BD_Académicos_JCE": ("JCE", {
            "academicos_jce": ("por institucion", "total general",
                               "Suma de jornadas completas equivalentes del personal académico"),
            "academicos_jce_doctorado": ("por nivel de formacion", "doctor",
                                         "JCE del personal cuyo grado máximo es doctor"),
            "academicos_jce_extranjeros": ("por nacionalidad", "extranjero",
                                           "JCE del personal de nacionalidad distinta de la chilena"),
        }),
    }

    filas_salida: list[dict] = []
    sin_alias: set[str] = set()

    for hoja, (unidad, variables) in HOJAS.items():
        if hoja not in libro.sheetnames:
            print(f"    [aviso] falta la hoja {hoja}")
            continue
        ws = libro[hoja]
        columnas = _columnas_compuestas(ws)
        indices = {v: _indice(columnas, g, h) for v, (g, h, _) in variables.items()}
        faltan = [v for v, i in indices.items() if i is None]
        if faltan:
            print(f"    [aviso] en {hoja} no se hallaron: {faltan}")

        for fila in ws.iter_rows(min_row=4, values_only=True):
            periodo = str(fila[0] or "")
            m = re.search(r"(20\d\d)", periodo)
            if not m:
                continue
            anio = int(m.group(1))
            if anio < DESDE:
                continue
            universidad = resolver(fila[2])
            if universidad is None:
                sin_alias.add(str(fila[2] or "").strip())
                continue

            for variable, indice in indices.items():
                if indice is None or indice >= len(fila):
                    continue
                valor = numero(fila[indice])
                if valor is None:
                    continue
                filas_salida.append(nav.dato(
                    FUENTE, universidad, variable, round(valor, 2), anio_dato=anio,
                    unidad=unidad, definicion=variables[variable][2],
                    url=BASES["personal_academico"][0]))

    libro.close()
    return filas_salida, sin_alias


# --- Titulados ----------------------------------------------------------------

# Categorías de `CARRERA CLASIFICACIÓN NIVEL 1` que componen cada variable. THE
# necesita doctorados otorgados y licenciaturas otorgadas por separado: la razón
# entre ambos es su métrica «Doctorate/bachelor ratio».
NIVELES = {
    "titulados_pregrado": ({"profesional con licenciatura", "licenciatura no conducente a titulo"},
                           "Titulaciones de pregrado con grado de licenciado"),
    "titulados_pregrado_total": (None,  # todo lo que la base clasifica como pregrado
                                 "Todas las titulaciones de pregrado, con y sin licenciatura"),
    "graduados_magister": ({"magister"}, "Graduados de programas de magíster"),
    "graduados_doctorado": ({"doctorado"}, "Graduados de programas de doctorado"),
    "titulados_especialidad_medica": ({"especialidad medica u odontologica"},
                                      "Titulados de especialidades médicas u odontológicas"),
}


def titulados(ruta: Path, resolver) -> tuple[list[dict], set[str]]:
    """Titulaciones por institución, año y nivel.

    La base trae una fila por programa y año con el total de titulaciones, de modo
    que hay que sumar. Son 234.000 filas: se recorre una sola vez, acumulando.
    """
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = libro[libro.sheetnames[0]]
    iterador = ws.iter_rows(values_only=True)
    encabezados = [plano(c) for c in next(iterador)]

    c_anio = encabezados.index("ano")
    c_total = encabezados.index("total titulaciones")
    c_inst = encabezados.index("nombre institucion")
    c_global = encabezados.index("nivel global")
    c_nivel = encabezados.index("carrera clasificacion nivel 1")

    acumulado: dict[tuple[str, int, str], float] = {}
    sin_alias: set[str] = set()

    for fila in iterador:
        m = re.search(r"(20\d\d)", str(fila[c_anio] or ""))
        if not m:
            continue
        anio = int(m.group(1))
        if anio < DESDE:
            continue
        universidad = resolver(fila[c_inst])
        if universidad is None:
            sin_alias.add(str(fila[c_inst] or "").strip())
            continue
        total = numero(fila[c_total])
        if not total:
            continue

        nivel = plano(fila[c_nivel])
        global_ = plano(fila[c_global])
        for variable, (categorias, _) in NIVELES.items():
            if categorias is None:
                if variable == "titulados_pregrado_total" and global_ == "pregrado":
                    acumulado[(universidad, anio, variable)] =                         acumulado.get((universidad, anio, variable), 0) + total
            elif nivel in categorias:
                acumulado[(universidad, anio, variable)] =                     acumulado.get((universidad, anio, variable), 0) + total

    libro.close()
    return [nav.dato(FUENTE, universidad, variable, round(valor, 0), anio_dato=anio,
                     unidad="titulaciones", definicion=NIVELES[variable][1],
                     url=BASES["titulados"][0])
            for (universidad, anio, variable), valor in sorted(acumulado.items())], sin_alias


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--solo-procesar", action="store_true")
    args = p.parse_args()

    if not args.solo_procesar:
        print("Descargando las bases del SIES…")
        descargar()

    base = nav.carpeta(FUENTE) / "raw"
    resolver = nav.resolvedor()
    filas: list[dict] = []
    sin_alias: set[str] = set()

    ruta = base / BASES["personal_academico"][1]
    if ruta.exists():
        print("\nPersonal académico:")
        nuevas, avisos = personal_academico(ruta, resolver)
        filas += nuevas
        sin_alias |= avisos
        universidades = {f["universidad"] for f in nuevas}
        anios = {f["anio_dato"] for f in nuevas}
        print(f"    {len(nuevas):,} datos · {len(universidades)} universidades · "
              f"{min(anios)}-{max(anios)}".replace(",", "."))

    ruta = base / BASES["titulados"][1]
    if ruta.exists():
        print("\nTitulados:")
        nuevas, avisos = titulados(ruta, resolver)
        filas += nuevas
        sin_alias |= avisos
        universidades = {f["universidad"] for f in nuevas}
        anios = {f["anio_dato"] for f in nuevas}
        print(f"    {len(nuevas):,} datos · {len(universidades)} universidades · "
              f"{min(anios)}-{max(anios)}".replace(",", "."))

    # La matrícula queda pendiente: el SIES la publica en un ZIP que contiene un
    # RAR, y este equipo no tiene con qué abrirlo (ni 7-Zip, ni WinRAR, ni unrar
    # en WSL). Se documenta en vez de instalarlo por cuenta propia.
    if (base / BASES["matricula"][1]).exists():
        print("\nMatrícula: descargada, sin procesar (viene en RAR; ver README).")

    salida = nav.escribir_procesado(FUENTE, filas)
    print(f"\n{len(filas):,} filas -> {salida}".replace(",", "."))
    if sin_alias:
        # La base del SIES cubre todas las instituciones de educación superior:
        # institutos profesionales y centros de formación técnica incluidos. Que
        # no tengan alias es lo esperado, no un fallo.
        print(f"  ({len(sin_alias)} instituciones fuera de la tabla `universidad`: "
              "institutos profesionales y CFT)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
