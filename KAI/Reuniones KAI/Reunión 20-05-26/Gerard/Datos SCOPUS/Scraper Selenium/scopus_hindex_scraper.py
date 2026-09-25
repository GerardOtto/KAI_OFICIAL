"""
scopus_hindex_scraper.py
========================
Extrae el h-index de cada autor en Scopus y lo escribe en la columna h_index
del archivo Excel. Configurado para OperaGX usando ChromeDriver 147.

REQUISITOS
----------
    pip install selenium openpyxl

    ChromeDriver 147 — descárgalo desde:
    https://googlechromelabs.github.io/chrome-for-testing/
    → Busca la versión 147 → descarga "chromedriver" para tu OS (win64 / mac-x64 / mac-arm64 / linux64)

USO
---
    1. Edita las variables OPERA_GX_PATH y CHROMEDRIVER_PATH más abajo.
    2. Pon este script en la misma carpeta que PUCV_Hindex_Scopus.xlsx
    3. Ejecuta:  python scopus_hindex_scraper.py

El script guarda el Excel después de cada autor. Si se interrumpe,
al volver a ejecutar salta los que ya tienen h-index.
"""

import time
import re
import sys
from pathlib import Path

# ──────────────────────────────────────────────
# ★  EDITA ESTAS DOS RUTAS  ★
# ──────────────────────────────────────────────

# Ruta al ejecutable de OperaGX:
#   Windows ejemplo: r"C:\Users\TU_USUARIO\AppData\Local\Programs\Opera GX\opera.exe"
#   macOS   ejemplo: "/Applications/Opera GX.app/Contents/MacOS/Opera"
OPERA_GX_PATH = r"C:\Users\UsuarioCompuElite\AppData\Local\Programs\Opera GX\opera.exe"

# Ruta al chromedriver.exe que descargaste:
#   Windows ejemplo: r"C:\tools\chromedriver\chromedriver.exe"
#   macOS   ejemplo: "/usr/local/bin/chromedriver"
CHROMEDRIVER_PATH = r"C:\Users\UsuarioCompuElite\Downloads\chromedriver-win64\chromedriver-win64\chromedriver.exe"

# ──────────────────────────────────────────────
# CONFIGURACIÓN DEL EXCEL
# ──────────────────────────────────────────────

EXCEL_PATH   = Path("PUCV_Hindex_Scopus.xlsx")
SHEET_NAME   = "Sheet1"
AUTHOR_COL   = 2    # columna B — Auth-ID
HINDEX_COL   = 3    # columna C — h_index
HEADER_ROW   = 1
FIRST_DATA   = 2

BASE_URL       = "https://www.scopus.com/authid/detail.uri?authorId={}"
PAGE_TIMEOUT   = 30   # segundos esperando que cargue la página
DELAY_BETWEEN  = 2    # segundos de pausa entre autores

# ──────────────────────────────────────────────
# IMPORTS
# ──────────────────────────────────────────────

try:
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.common.exceptions import TimeoutException, NoSuchElementException
except ImportError:
    sys.exit("❌  Instala selenium:  pip install selenium")

try:
    import openpyxl
except ImportError:
    sys.exit("❌  Instala openpyxl:  pip install openpyxl")

# ──────────────────────────────────────────────
# INICIALIZAR DRIVER
# ──────────────────────────────────────────────

def init_driver() -> webdriver.Chrome:

    opts = Options()
    opts.binary_location = OPERA_GX_PATH

    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--window-size=1280,900")

    # anti-detección
    opts.add_argument("--disable-blink-features=AutomationControlled")

    opts.add_experimental_option(
        "excludeSwitches",
        ["enable-automation"]
    )

    opts.add_experimental_option(
        "useAutomationExtension",
        False
    )

    service = Service(executable_path=CHROMEDRIVER_PATH)

    driver = webdriver.Chrome(
        service=service,
        options=opts
    )

    driver.execute_script(
        "Object.defineProperty(navigator, 'webdriver', "
        "{get: () => undefined})"
    )

    return driver


# ──────────────────────────────────────────────
# EXTRACCIÓN DEL H-INDEX
# ──────────────────────────────────────────────

def get_hindex(driver: webdriver.Chrome, author_id: str) -> str | None:
    """
    Navega al perfil del autor en Scopus y extrae el h-index.
    El h-index corresponde al TERCER span:
    <span data-testid="unclickable-count">...</span>
    """

    url = BASE_URL.format(author_id.strip())
    driver.get(url)

    wait = WebDriverWait(driver, PAGE_TIMEOUT)

    try:
        # Esperar a que aparezcan las métricas
        wait.until(
            EC.presence_of_all_elements_located(
                (
                    By.CSS_SELECTOR,
                    "span[data-testid='unclickable-count']"
                )
            )
        )

        # Espera extra para render dinámico React
        time.sleep(2)

        counts = driver.find_elements(
            By.CSS_SELECTOR,
            "span[data-testid='unclickable-count']"
        )

        # DEBUG opcional
        # print(f"DEBUG counts encontrados: {len(counts)}")
        # for i, el in enumerate(counts):
        #     print(f"{i}: {el.text}")

        # El h-index es el TERCER elemento
        if len(counts) >= 3:

            hindex = counts[2].text.strip()

            if re.match(r"^\d+$", hindex):
                return hindex

    except TimeoutException:
        print("⚠️ Timeout esperando métricas")

    except Exception as e:
        print(f"⚠️ Error scraping h-index: {e}")

    return None
# ──────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────

def main():
    # Validar rutas
    if not EXCEL_PATH.exists():
        sys.exit(
            f"❌  No se encontró el archivo Excel: {EXCEL_PATH.resolve()}\n"
            f"    Asegúrate de ejecutar el script desde la misma carpeta que el .xlsx\n"
            f"    o edita la variable EXCEL_PATH al inicio del script."
        )
    if not Path(OPERA_GX_PATH).exists():
        sys.exit(
            f"❌  No se encontró OperaGX en:\n    {OPERA_GX_PATH}\n"
            f"    Edita la variable OPERA_GX_PATH al inicio del script."
        )
    if not Path(CHROMEDRIVER_PATH).exists():
        sys.exit(
            f"❌  No se encontró chromedriver en:\n    {CHROMEDRIVER_PATH}\n"
            f"    Descárgalo desde: https://googlechromelabs.github.io/chrome-for-testing/\n"
            f"    y edita la variable CHROMEDRIVER_PATH al inicio del script."
        )

    # Cargar Excel
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    max_row = ws.max_row

    # Recopilar filas pendientes (sin h-index)
    pending = [
        (row, str(ws.cell(row=row, column=AUTHOR_COL).value))
        for row in range(FIRST_DATA, max_row + 1)
        if ws.cell(row=row, column=AUTHOR_COL).value
        and not str(ws.cell(row=row, column=HINDEX_COL).value or "").strip()
    ]

    total = max_row - HEADER_ROW
    done  = total - len(pending)
    print(f"📊  Autores totales:   {total}")
    print(f"✅  Ya procesados:    {done}")
    print(f"🔄  Por procesar:     {len(pending)}\n")

    if not pending:
        print("¡Todo listo! No hay autores pendientes.")
        return

    print("🌐  Iniciando OperaGX…  (NO lo cierres mientras corre el script)\n")
    print("    ⚠️   Si Scopus pide login institucional (PUCV), hazlo manualmente")
    print("        en la ventana del navegador y luego presiona ENTER aquí.\n")

    driver = init_driver()
    driver.get("https://www.scopus.com")
    input("    Presiona ENTER cuando estés listo para comenzar... ")

    errors    = []
    not_found = []

    for i, (row, author_id) in enumerate(pending, start=1):
        author_name = ws.cell(row=row, column=1).value or ""
        print(f"[{i:>4}/{len(pending)}] {author_name:<40} ID: {author_id}", end="  →  ", flush=True)

        try:
            hindex = get_hindex(driver, author_id)
        except Exception as e:
            print(f"❌ ERROR: {e}")
            errors.append((row, author_id, str(e)))
            time.sleep(DELAY_BETWEEN)
            continue

        if hindex is not None:
            print(f"h-index = {hindex}")
            ws.cell(row=row, column=HINDEX_COL).value = int(hindex)
        else:
            print("⚠️  no encontrado")
            not_found.append((row, author_id, author_name))
            ws.cell(row=row, column=HINDEX_COL).value = "N/A"

        wb.save(EXCEL_PATH)
        time.sleep(DELAY_BETWEEN)

    driver.quit()

    # Resumen
    print("\n" + "=" * 60)
    print("RESUMEN FINAL")
    print("=" * 60)
    print(f"  ✅ Procesados exitosamente: {len(pending) - len(errors) - len(not_found)}")
    print(f"  ⚠️  H-index no encontrado:  {len(not_found)}")
    print(f"  ❌ Errores:                 {len(errors)}")

    if not_found:
        print("\n  Autores sin h-index (perfil vacío o sin métricas):")
        for row, aid, name in not_found:
            print(f"    Fila {row:>4} | {aid} | {name}")

    if errors:
        print("\n  Errores de conexión/scraping:")
        for row, aid, err in errors:
            print(f"    Fila {row:>4} | {aid} | {err}")

    print(f"\n💾  Resultado guardado en: {EXCEL_PATH.resolve()}")


if __name__ == "__main__":
    main()